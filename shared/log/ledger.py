# -*- coding: utf-8 -*-
"""
Aminsa Performance Ledger & Audit Logger
Logs cycles to reports/daily_performance.xlsx and auto-maintains "Last 24 h" summary sheet.
"""

import os
import sys
import json
import time
from datetime import datetime, timedelta

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
REPORTS_DIR = os.path.join(ROOT_DIR, "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)
XLSX_PATH = os.path.join(REPORTS_DIR, "daily_performance.xlsx")

def log_row(slot, cycle_type="IMAGE", topic="General", face_score=1.0, render_s=0.0, post_ids=None, status="POSTED", cycle_id=None):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        if os.path.exists(XLSX_PATH):
            wb = openpyxl.load_workbook(XLSX_PATH)
        else:
            wb = openpyxl.Workbook()

        today_str = datetime.now().strftime("%Y%m%d")
        sheet_name = f"Log_{today_str}"

        # 1. Append to Today's Detailed Sheet
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            headers = ["Cycle ID", "Timestamp", "Type", "Slot", "Topic", "Face Score", "Render Sec", "Post IDs", "Status"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0F172A")
                cell.alignment = Alignment(horizontal="center")
        else:
            ws = wb[sheet_name]

        cid = cycle_id or f"CYC-{cycle_type}-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
        ts_now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        ws.append([
            cid,
            ts_now,
            cycle_type,
            slot,
            topic,
            round(face_score, 4),
            round(render_s, 2),
            json.dumps(post_ids) if isinstance(post_ids, (dict, list)) else str(post_ids or {}),
            status
        ])

        # 2. Update / Rebuild "Last 24 h" Rolling Sheet
        summary_sheet_name = "Last 24 h"
        if summary_sheet_name in wb.sheetnames:
            del wb[summary_sheet_name]

        ws_sum = wb.create_sheet(summary_sheet_name, 0) # Create as first tab
        headers_sum = ["Cycle ID", "Timestamp", "Type", "Slot", "Topic", "Render Sec", "Status"]
        for col, h in enumerate(headers_sum, 1):
            cell = ws_sum.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E3A8A")
            cell.alignment = Alignment(horizontal="center")

        cutoff = datetime.now() - timedelta(hours=24)
        for s in wb.sheetnames:
            if s.startswith("Log_"):
                sheet = wb[s]
                for r in sheet.iter_rows(min_row=2, values_only=True):
                    if r and r[1]:
                        try:
                            row_time = datetime.strptime(str(r[1]), "%Y-%m-%d %H:%M:%S")
                            if row_time >= cutoff:
                                ws_sum.append([r[0], r[1], r[2], r[3], r[4], r[6], r[8]])
                        except Exception:
                            pass

        wb.save(XLSX_PATH)
        print(f"[+] Ledger updated: {cid} -> {XLSX_PATH}")
        return True
    except Exception as e:
        print(f"[WARN] Ledger logging error: {e}")
        return False
